import shutil
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

from run_production import run
import transform

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"


def _make_drifted_input(tmp_path: Path) -> Path:
    source = SAMPLES / "input.xlsx"
    drifted = tmp_path / "input_drifted.xlsx"
    shutil.copy2(source, drifted)

    wb = load_workbook(drifted)
    ws = wb[transform.SHEET_NAME]
    ws["BW9"] = "drift"
    wb.save(drifted)

    return drifted


def _make_resized_input(tmp_path: Path, target_last_row: int) -> Path:
    """Copy the sample and resize its data to target_last_row rows.

    The structure (headers, columns, formats) is unchanged; only the number
    of data rows and the row-7 SUBTOTAL ranges differ, mimicking an incoming
    file of a different size.
    """
    source = SAMPLES / "input.xlsx"
    resized = tmp_path / f"input_resized_{target_last_row}.xlsx"
    shutil.copy2(source, resized)

    wb = load_workbook(resized)
    ws = wb[transform.SHEET_NAME]
    current_last = transform.data_row(ws)
    max_col = ws.max_column

    if target_last_row > current_last:
        for new_row in range(current_last + 1, target_last_row + 1):
            for col in range(1, max_col + 1):
                src = ws.cell(current_last, col)
                dst = ws.cell(new_row, col)
                dst.value = src.value
                dst._style = copy(src._style)
    elif target_last_row < current_last:
        ws.delete_rows(target_last_row + 1, current_last - target_last_row)

    for col in range(1, max_col + 1):
        cell = ws.cell(7, col)
        value = cell.value
        if isinstance(value, str) and value.startswith("="):
            match = transform._ROW7_SUBTOTAL_FULL_RE.match(value)
            if match:
                cell.value = (
                    f"=SUBTOTAL({match.group(1)},{match.group(2)}10:"
                    f"{match.group(3)}{target_last_row + 1})"
                )

    wb.save(resized)
    return resized


def test_clean_input_gets_transformed(tmp_path):
    output_path = tmp_path / "actual_output.xlsx"
    needs_review_dir = tmp_path / "needs_review"

    exit_code = run(
        input_path=SAMPLES / "input.xlsx",
        output_path=output_path,
        sample_path=SAMPLES / "input.xlsx",
        report_dir=tmp_path / "reports",
        needs_review_dir=needs_review_dir,
    )

    assert exit_code == 0
    assert output_path.exists()
    assert not needs_review_dir.exists()


def test_drifted_input_is_quarantined_not_transformed(tmp_path):
    drifted_input = _make_drifted_input(tmp_path)
    output_path = tmp_path / "actual_output.xlsx"
    needs_review_dir = tmp_path / "needs_review"

    exit_code = run(
        input_path=drifted_input,
        output_path=output_path,
        sample_path=SAMPLES / "input.xlsx",
        report_dir=tmp_path / "reports",
        needs_review_dir=needs_review_dir,
    )

    assert exit_code == 1
    assert not output_path.exists()

    case_dirs = list(needs_review_dir.iterdir())
    assert len(case_dirs) == 1
    quarantined = case_dirs[0]
    assert (quarantined / "input_drifted.xlsx").exists()
    assert (quarantined / "input_vs_sample.html").exists()


def test_larger_input_is_transformed_not_quarantined(tmp_path):
    bigger = _make_resized_input(tmp_path, target_last_row=200)
    output_path = tmp_path / "actual_output.xlsx"
    needs_review_dir = tmp_path / "needs_review"

    exit_code = run(
        input_path=bigger,
        output_path=output_path,
        sample_path=SAMPLES / "input.xlsx",
        report_dir=tmp_path / "reports",
        needs_review_dir=needs_review_dir,
    )

    assert exit_code == 0
    assert output_path.exists()
    assert not needs_review_dir.exists()
    ws = load_workbook(output_path)[transform.SHEET_NAME]
    assert transform.data_row(ws) == 200


def test_smaller_input_is_transformed_not_quarantined(tmp_path):
    smaller = _make_resized_input(tmp_path, target_last_row=50)
    output_path = tmp_path / "actual_output.xlsx"
    needs_review_dir = tmp_path / "needs_review"

    exit_code = run(
        input_path=smaller,
        output_path=output_path,
        sample_path=SAMPLES / "input.xlsx",
        report_dir=tmp_path / "reports",
        needs_review_dir=needs_review_dir,
    )

    assert exit_code == 0
    assert output_path.exists()
    assert not needs_review_dir.exists()
    ws = load_workbook(output_path)[transform.SHEET_NAME]
    assert transform.data_row(ws) == 50
