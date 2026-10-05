import hashlib
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

import transform
from run_batch import discover_workbooks, run_batch

ROOT = Path(__file__).resolve().parents[1]
SAMPLE_INPUT = ROOT / "samples" / "input.xlsx"


def _require_sample() -> None:
    if not SAMPLE_INPUT.exists():
        pytest.skip("samples/input.xlsx is required for batch tests")


def _copy_sample(input_dir: Path, name: str) -> Path:
    input_dir.mkdir(parents=True, exist_ok=True)
    target = input_dir / name
    shutil.copy2(SAMPLE_INPUT, target)
    return target


def _drift(path: Path) -> None:
    wb = load_workbook(path)
    wb[transform.SHEET_NAME]["BW9"] = "drift"
    wb.save(path)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture(scope="module")
def basic_batch(tmp_path_factory):
    _require_sample()
    base = tmp_path_factory.mktemp("batch_basic")
    input_dir = base / "input"
    output_dir = base / "output"
    report_dir = base / "reports"
    needs_review_dir = base / "needs_review"

    _copy_sample(input_dir, "Alpha.xlsx")
    _copy_sample(input_dir, "Beta.xlsx")
    (input_dir / "~$Alpha.xlsx").write_text("temporary")
    (input_dir / "notes.txt").write_text("ignore")
    (input_dir / "data.csv").write_text("ignore")
    (input_dir / "nested").mkdir()
    (input_dir / "nested" / "Nested.xlsx").write_bytes(b"ignored")

    before = {path.name: _sha256(path) for path in discover_workbooks(input_dir)}
    code = run_batch(
        input_dir,
        output_dir,
        sample_path=SAMPLE_INPUT,
        report_dir=report_dir,
        needs_review_dir=needs_review_dir,
    )
    after = {path.name: _sha256(path) for path in discover_workbooks(input_dir)}

    return {
        "code": code,
        "input_dir": input_dir,
        "output_dir": output_dir,
        "report_dir": report_dir,
        "needs_review_dir": needs_review_dir,
        "before": before,
        "after": after,
    }


@pytest.fixture(scope="module")
def failure_batch(tmp_path_factory):
    _require_sample()
    base = tmp_path_factory.mktemp("batch_failure")
    input_dir = base / "input"
    output_dir = base / "output"
    report_dir = base / "reports"
    needs_review_dir = base / "needs_review"

    _copy_sample(input_dir, "Good.xlsx")
    bad_path = _copy_sample(input_dir, "Bad.xlsx")
    _drift(bad_path)
    (input_dir / "Broken.xlsx").write_bytes(b"not a workbook")

    code = run_batch(
        input_dir,
        output_dir,
        sample_path=SAMPLE_INPUT,
        report_dir=report_dir,
        needs_review_dir=needs_review_dir,
    )

    return {
        "code": code,
        "input_dir": input_dir,
        "output_dir": output_dir,
        "report_dir": report_dir,
        "needs_review_dir": needs_review_dir,
    }


@pytest.fixture(scope="module")
def existing_output_batch(tmp_path_factory):
    _require_sample()
    base = tmp_path_factory.mktemp("batch_existing_output")
    input_dir = base / "input"
    output_dir = base / "output"
    report_dir = base / "reports"
    needs_review_dir = base / "needs_review"

    _copy_sample(input_dir, "Existing.xlsx")
    _copy_sample(input_dir, "Good.xlsx")
    output_dir.mkdir(parents=True)
    sentinel = output_dir / "Existing.xlsx"
    sentinel.write_text("sentinel")

    code = run_batch(
        input_dir,
        output_dir,
        sample_path=SAMPLE_INPUT,
        report_dir=report_dir,
        needs_review_dir=needs_review_dir,
        overwrite=False,
    )

    return {
        "code": code,
        "input_dir": input_dir,
        "output_dir": output_dir,
        "sentinel": sentinel,
    }


def test_output_filenames_exactly_match_input(basic_batch):
    output_dir = basic_batch["output_dir"]
    assert sorted(path.name for path in output_dir.iterdir()) == [
        "Alpha.xlsx",
        "Beta.xlsx",
    ]


def test_multiple_files_transform_successfully(basic_batch):
    assert basic_batch["code"] == 0
    for name in ("Alpha.xlsx", "Beta.xlsx"):
        ws = load_workbook(basic_batch["output_dir"] / name)[transform.SHEET_NAME]
        assert ws.max_column == transform.OUTPUT_MAX_COL


def test_excel_temp_files_are_ignored(basic_batch):
    output_names = {path.name for path in basic_batch["output_dir"].iterdir()}
    assert not any(name.startswith("~$") for name in output_names)


def test_non_excel_files_are_ignored(basic_batch):
    output_names = {path.name for path in basic_batch["output_dir"].iterdir()}
    assert "notes.txt" not in output_names
    assert "data.csv" not in output_names


def test_subdirectories_are_not_recursed(basic_batch):
    output_names = {path.name for path in basic_batch["output_dir"].iterdir()}
    assert "Nested.xlsx" not in output_names


def test_source_files_are_not_modified(basic_batch):
    assert basic_batch["before"] == basic_batch["after"]


def test_unique_report_directory_per_file(basic_batch):
    assert (basic_batch["report_dir"] / "Alpha" / "input_vs_sample.json").exists()
    assert (basic_batch["report_dir"] / "Beta" / "input_vs_sample.json").exists()


def test_rejected_file_does_not_stop_other_files(failure_batch):
    assert failure_batch["code"] == 1
    assert (failure_batch["output_dir"] / "Good.xlsx").exists()
    assert not (failure_batch["output_dir"] / "Bad.xlsx").exists()
    assert not (failure_batch["output_dir"] / "Broken.xlsx").exists()


def test_rejected_file_is_quarantined(failure_batch):
    bad_needs_review = failure_batch["needs_review_dir"] / "Bad"
    assert list(bad_needs_review.rglob("Bad.xlsx"))
    assert list(bad_needs_review.rglob("input_vs_sample.html"))


def test_unexpected_file_error_does_not_stop_other_files(failure_batch):
    assert (failure_batch["output_dir"] / "Good.xlsx").exists()
    assert not (failure_batch["output_dir"] / "Broken.xlsx").exists()


def test_existing_output_is_preserved_without_overwrite(existing_output_batch):
    assert existing_output_batch["code"] == 1
    assert existing_output_batch["sentinel"].read_text() == "sentinel"
    assert (existing_output_batch["output_dir"] / "Good.xlsx").exists()


def test_overwrite_safely_replaces_existing_output(tmp_path):
    _require_sample()
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    _copy_sample(input_dir, "Report.xlsx")
    output_dir.mkdir()
    sentinel = output_dir / "Report.xlsx"
    sentinel.write_text("sentinel")

    code = run_batch(
        input_dir,
        output_dir,
        sample_path=SAMPLE_INPUT,
        report_dir=tmp_path / "reports",
        needs_review_dir=tmp_path / "needs_review",
        overwrite=True,
    )

    assert code == 0
    assert sentinel.read_bytes() != b"sentinel"
    ws = load_workbook(output_dir / "Report.xlsx")[transform.SHEET_NAME]
    assert ws.max_column == transform.OUTPUT_MAX_COL


def test_input_and_output_directories_cannot_be_the_same(tmp_path):
    input_dir = tmp_path / "same"
    input_dir.mkdir()

    code = run_batch(
        input_dir,
        input_dir,
        sample_path=SAMPLE_INPUT,
        report_dir=tmp_path / "reports",
        needs_review_dir=tmp_path / "needs_review",
    )

    assert code == 1


def test_equivalent_resolved_directory_paths_are_rejected(tmp_path):
    input_dir = tmp_path / "same"
    input_dir.mkdir()
    output_dir = input_dir / "."

    code = run_batch(
        input_dir,
        output_dir,
        sample_path=SAMPLE_INPUT,
        report_dir=tmp_path / "reports",
        needs_review_dir=tmp_path / "needs_review",
    )

    assert code == 1


def test_empty_input_directory_returns_success(tmp_path):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    output_dir = tmp_path / "output"

    code = run_batch(
        input_dir,
        output_dir,
        sample_path=SAMPLE_INPUT,
        report_dir=tmp_path / "reports",
        needs_review_dir=tmp_path / "needs_review",
    )

    assert code == 0
    assert not output_dir.exists()


def test_discover_workbooks_uses_deterministic_alphabetical_order(tmp_path):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    (input_dir / "Zebra.xlsx").write_bytes(b"xlsx")
    (input_dir / "apple.xlsx").write_bytes(b"xlsx")
    (input_dir / "~$apple.xlsx").write_bytes(b"xlsx")
    (input_dir / "notes.txt").write_bytes(b"text")
    (input_dir / "nested").mkdir()
    (input_dir / "nested" / "Nested.xlsx").write_bytes(b"xlsx")

    discovered = [path.name for path in discover_workbooks(input_dir)]

    assert discovered == ["apple.xlsx", "Zebra.xlsx"]
