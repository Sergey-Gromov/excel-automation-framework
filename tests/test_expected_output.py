from openpyxl import load_workbook

import transform


def test_actual_applies_intended_changes(paths):
    if not paths["input"].exists() or not paths["actual"].exists():
        return

    input_ws = load_workbook(paths["input"], data_only=False)[transform.SHEET_NAME]
    actual_ws = load_workbook(paths["actual"], data_only=False)[transform.SHEET_NAME]

    assert actual_ws.max_column == transform.OUTPUT_MAX_COL

    last_data_row = transform.data_row(input_ws)

    def subtotal(col: str) -> str:
        return f"=SUBTOTAL(109,{col}10:{col}{last_data_row + 1})"

    assert actual_ws["BW9"].value == input_ws["FF9"].value
    assert actual_ws["BW9"].value.strip() == "Сумма годовой корректировки"
    assert actual_ws["BX9"].value == input_ws["FG9"].value
    assert actual_ws["BX9"].value.strip() == "Перерасчет суммы годовой корректировки"
    assert actual_ws["BY9"].value == input_ws["BW9"].value
    assert actual_ws["BY9"].value.strip() == "Оплачено нарастающим итогом с начала периода"
    # Identity block: final EC (133) <- source EC (133), position unchanged.
    assert actual_ws["EC9"].value == input_ws["EC9"].value
    assert actual_ws["EC9"].value.strip() == "Соглашение о реструктуризации"

    assert actual_ws["BV7"].value == subtotal("BV")
    assert actual_ws["BW7"].value == subtotal("BW")
    assert actual_ws["BX7"].value == subtotal("BX")
    assert actual_ws["BY7"].value == subtotal("BY")
    # Identity-block boundary: final EC (133) <- source EC (133), translated.
    assert actual_ws["EC7"].value == subtotal("EC")
    # Shifted block straddling the deletion: final EA (131) <- source DY (129).
    assert actual_ws["EA7"].value == subtotal("EA")
    assert actual_ws["BB7"].value == subtotal("BB")
    assert actual_ws["CV7"].value == subtotal("CV")
    assert actual_ws["BB7"].number_format == transform.TOTAL_NUMBER_FORMAT
    assert actual_ws["CV7"].number_format == transform.TOTAL_NUMBER_FORMAT

    for col in ("I", "J", "S", "T"):
        for row in range(10, last_data_row + 1):
            coord = f"{col}{row}"
            assert actual_ws[coord].number_format == transform.DATE_NUMBER_FORMAT
            assert actual_ws[coord].value == input_ws[coord].value

    assert actual_ws["AE9"].value == '="Июль 2023"'
    assert actual_ws["BQ9"].value == '="Июнь 2026"'

    moved_values = (
        ("BW10", "FF10"),   # moved: final BW (75) <- source FF (162)
        ("BX10", "FG10"),   # moved: final BX (76) <- source FG (163)
        ("BY10", "BW10"),   # shifted: final BY (77) <- source BW (75)
        ("EA12", "DY12"),   # shifted, straddles deletion: final EA (131) <- source DY (129)
        ("EC12", "EC12"),   # identity block: final EC (133) <- source EC (133)
    )
    for actual_coord, input_coord in moved_values:
        assert actual_ws[actual_coord].value == input_ws[input_coord].value

    assert sorted(str(rng) for rng in actual_ws.merged_cells.ranges) == [
        "DH8:EM8",
        "EN8:EP8",
    ]
