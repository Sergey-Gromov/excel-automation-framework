from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

import transform


def _style_signature(style):
    return ElementTree.tostring(style.to_tree(), encoding="unicode")


def _remapped_merged_ranges(ws):
    remapped = []
    for rng in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(rng))
        remapped.append(
            f"{get_column_letter(transform.final_for_source_column(min_col))}{min_row}:"
            f"{get_column_letter(transform.final_for_source_column(max_col))}{max_row}"
        )
    return sorted(remapped)


def test_no_unexpected_changes_from_input(paths):
    if not paths["input"].exists() or not paths["actual"].exists():
        return

    input_wb = load_workbook(paths["input"], data_only=False)
    actual_wb = load_workbook(paths["actual"], data_only=False)

    assert input_wb.sheetnames == actual_wb.sheetnames

    input_ws = input_wb[transform.SHEET_NAME]
    actual_ws = actual_wb[transform.SHEET_NAME]

    assert actual_ws.max_row == transform.EXPECTED_MAX_ROW
    assert actual_ws.max_column == transform.EXPECTED_MAX_COL

    date_cols = set(transform.DATE_COLUMNS.values())
    new_formula_coords = set(transform.NEW_FORMULAS)
    style_attrs = ("font", "fill", "border", "alignment", "protection")

    for row in range(1, transform.EXPECTED_MAX_ROW + 1):
        for final_col in range(1, transform.EXPECTED_MAX_COL + 1):
            source_col = transform.source_for_final_column(final_col)
            source_cell = input_ws.cell(row, source_col)
            actual_cell = actual_ws.cell(row, final_col)

            if actual_cell.coordinate in transform.NEW_FORMULAS:
                expected_value = transform.NEW_FORMULAS[actual_cell.coordinate]
            else:
                expected_value = source_cell.value
                if isinstance(expected_value, str) and expected_value.startswith("="):
                    expected_value = transform.translate_formula(expected_value)

            assert actual_cell.value == expected_value, (
                f"Value mismatch at {actual_cell.coordinate}: "
                f"expected {expected_value!r}, actual {actual_cell.value!r}"
            )

            expected_format = source_cell.number_format
            if actual_cell.coordinate in new_formula_coords:
                expected_format = transform.TOTAL_NUMBER_FORMAT
            elif row >= 10 and final_col in date_cols:
                expected_format = transform.DATE_NUMBER_FORMAT

            assert actual_cell.number_format == expected_format, (
                f"Number format mismatch at {actual_cell.coordinate}: "
                f"expected {expected_format!r}, actual {actual_cell.number_format!r}"
            )

            for attr in style_attrs:
                assert _style_signature(getattr(actual_cell, attr)) == _style_signature(
                    getattr(source_cell, attr)
                ), f"{attr} mismatch at {actual_cell.coordinate}"

            assert (actual_cell.hyperlink is None) == (source_cell.hyperlink is None), (
                f"Hyperlink presence mismatch at {actual_cell.coordinate}"
            )
            assert (actual_cell.comment is None) == (source_cell.comment is None), (
                f"Comment presence mismatch at {actual_cell.coordinate}"
            )

    for row in range(1, transform.EXPECTED_MAX_ROW + 1):
        assert actual_ws.row_dimensions[row].height == input_ws.row_dimensions[row].height, (
            f"Row height mismatch at row {row}"
        )

    for final_col in range(1, transform.EXPECTED_MAX_COL + 1):
        source_col = transform.source_for_final_column(final_col)
        final_letter = get_column_letter(final_col)
        source_letter = get_column_letter(source_col)

        actual_dim = actual_ws.column_dimensions[final_letter]
        source_dim = input_ws.column_dimensions[source_letter]

        for attr in ("width", "bestFit", "hidden", "outlineLevel"):
            assert getattr(actual_dim, attr) == getattr(source_dim, attr), (
                f"Column {attr} mismatch at {final_letter}"
            )

    assert sorted(str(rng) for rng in actual_ws.merged_cells.ranges) == _remapped_merged_ranges(input_ws)

    assert actual_ws.freeze_panes == input_ws.freeze_panes
    assert list(actual_ws.tables) == list(input_ws.tables)
    assert len(actual_ws.data_validations.dataValidation) == len(
        input_ws.data_validations.dataValidation
    )
    assert len(list(actual_ws.conditional_formatting)) == len(list(input_ws.conditional_formatting))
