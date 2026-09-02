from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import transform
from run_production import run


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_INPUT = ROOT / "samples" / "input.xlsx"
DATE_SOURCE = ROOT / "sh513.xlsx"


def _require_samples() -> None:
    if not SAMPLE_INPUT.exists() or not DATE_SOURCE.exists():
        pytest.skip("Put samples/input.xlsx and sh513.xlsx into the repository first.")


def _make_date_source(path: Path, value) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = transform.SHEET_NAME
    ws["B1"] = value
    wb.save(path)


@pytest.fixture(scope="module")
def transformed_output(tmp_path_factory):
    _require_samples()
    tmp_dir = tmp_path_factory.mktemp("date_header_transform")
    output_path = tmp_dir / "actual_output.xlsx"
    transform.transform(SAMPLE_INPUT, output_path, DATE_SOURCE)
    return output_path


def _load_sheet(path: Path):
    return load_workbook(path, data_only=False)[transform.SHEET_NAME]


def test_target_header_dates_are_replaced(transformed_output):
    ws = _load_sheet(transformed_output)
    expected_date = transform.read_replacement_date(DATE_SOURCE)

    assert ws["CU9"].value == f"БУ ДЗ всего на {expected_date}"
    assert (
        ws["CV9"].value
        == f"Сумма просроченной ДЗ с учетом сроков оплаты счетов по состоянию на {expected_date} "
    )
    assert ws["CY9"].value == f"БУ Авансы на {expected_date}"


def test_unrelated_dated_headers_are_unchanged(transformed_output):
    input_ws = _load_sheet(SAMPLE_INPUT)
    actual_ws = _load_sheet(transformed_output)

    assert actual_ws["BZ9"].value == input_ws["BX9"].value
    assert actual_ws["CA9"].value == input_ws["BY9"].value


def test_normal_rows_are_preserved(transformed_output):
    input_ws = _load_sheet(SAMPLE_INPUT)
    actual_ws = _load_sheet(transformed_output)

    pairs = (
        ("A10", "A10"),
        ("BW10", "FF10"),
        ("BX10", "FG10"),
        ("BY10", "BW10"),
        ("CU10", "CS10"),
        ("CV10", "CT10"),
        ("CY10", "CW10"),
    )
    for actual_coord, input_coord in pairs:
        assert actual_ws[actual_coord].value == input_ws[input_coord].value


def test_formula_placement_is_preserved(transformed_output):
    ws = _load_sheet(transformed_output)

    assert ws["CU7"].value == "=SUBTOTAL(109,CU10:CU101)"
    assert ws["BY7"].value == "=SUBTOTAL(109,BY10:BY101)"
    assert ws["BB7"].value == transform.NEW_FORMULAS["BB7"]
    assert ws["CV7"].value == transform.NEW_FORMULAS["CV7"]
    assert ws["AE9"].value == '="Июль 2023"'


def test_blank_values_remain_blank(tmp_path):
    _require_samples()
    input_path = tmp_path / "input.xlsx"
    input_path.write_bytes(SAMPLE_INPUT.read_bytes())

    wb = load_workbook(input_path)
    ws = wb[transform.SHEET_NAME]
    ws["CS10"] = None
    ws["CT10"] = None
    ws["CW10"] = None
    wb.save(input_path)

    date_source = tmp_path / "date_source.xlsx"
    _make_date_source(date_source, datetime(2026, 6, 30))

    output_path = tmp_path / "actual_output.xlsx"
    transform.transform(input_path, output_path, date_source)

    ws = _load_sheet(output_path)
    assert ws["CU10"].value is None
    assert ws["CV10"].value is None
    assert ws["CY10"].value is None


def test_string_date_source_is_supported(tmp_path):
    _require_samples()
    input_path = tmp_path / "input.xlsx"
    input_path.write_bytes(SAMPLE_INPUT.read_bytes())

    date_source = tmp_path / "date_source.xlsx"
    _make_date_source(date_source, "15.07.2026")

    output_path = tmp_path / "actual_output.xlsx"
    transform.transform(input_path, output_path, date_source)

    ws = _load_sheet(output_path)
    assert ws["CU9"].value == "БУ ДЗ всего на 15.07.2026"
    assert ws["CV9"].value.endswith("на 15.07.2026 ")
    assert ws["CY9"].value == "БУ Авансы на 15.07.2026"


def test_missing_date_source_raises(tmp_path):
    _require_samples()
    with pytest.raises(ValueError, match="не найден"):
        transform.transform(
            SAMPLE_INPUT,
            tmp_path / "actual_output.xlsx",
            tmp_path / "missing_date_source.xlsx",
        )


def test_invalid_date_source_value_raises(tmp_path):
    _require_samples()
    date_source = tmp_path / "date_source.xlsx"
    _make_date_source(date_source, "not-a-date")

    with pytest.raises(ValueError, match="Не удалось распознать дату"):
        transform.transform(SAMPLE_INPUT, tmp_path / "actual_output.xlsx", date_source)


def test_missing_dated_header_raises(tmp_path):
    _require_samples()
    input_path = tmp_path / "input.xlsx"
    input_path.write_bytes(SAMPLE_INPUT.read_bytes())

    wb = load_workbook(input_path)
    wb[transform.SHEET_NAME]["CS9"] = "drift"
    wb.save(input_path)

    date_source = tmp_path / "date_source.xlsx"
    _make_date_source(date_source, datetime(2026, 6, 30))

    with pytest.raises(ValueError, match="БУ ДЗ всего на"):
        transform.transform(input_path, tmp_path / "actual_output.xlsx", date_source)


def test_duplicate_dated_header_raises(tmp_path):
    _require_samples()
    input_path = tmp_path / "input.xlsx"
    input_path.write_bytes(SAMPLE_INPUT.read_bytes())

    wb = load_workbook(input_path)
    wb[transform.SHEET_NAME]["CU9"] = "БУ Авансы на 31.10.2025"
    wb.save(input_path)

    date_source = tmp_path / "date_source.xlsx"
    _make_date_source(date_source, datetime(2026, 6, 30))

    with pytest.raises(ValueError, match="БУ Авансы на"):
        transform.transform(input_path, tmp_path / "actual_output.xlsx", date_source)


def test_date_only_header_change_passes_production_gate(tmp_path):
    _require_samples()
    incoming = tmp_path / "incoming.xlsx"
    incoming.write_bytes(SAMPLE_INPUT.read_bytes())

    wb = load_workbook(incoming)
    ws = wb[transform.SHEET_NAME]
    ws["CS9"] = "БУ ДЗ всего на 01.11.2025"
    ws["CT9"] = (
        "Сумма просроченной ДЗ с учетом сроков оплаты счетов "
        "по состоянию на 01.11.2025 "
    )
    ws["CW9"] = "БУ Авансы на 01.11.2025"
    wb.save(incoming)

    output_path = tmp_path / "actual_output.xlsx"
    exit_code = run(
        input_path=incoming,
        output_path=output_path,
        sample_path=SAMPLE_INPUT,
        report_dir=tmp_path / "reports",
        needs_review_dir=tmp_path / "needs_review",
    )

    assert exit_code == 0
    assert output_path.exists()
    assert not (tmp_path / "needs_review").exists()

    ws = _load_sheet(output_path)
    expected_date = transform.read_replacement_date(DATE_SOURCE)
    assert ws["CU9"].value == f"БУ ДЗ всего на {expected_date}"
    assert ws["CV9"].value.endswith(f"на {expected_date} ")
    assert ws["CY9"].value == f"БУ Авансы на {expected_date}"


def test_non_date_header_drift_still_quarantined(tmp_path):
    _require_samples()
    incoming = tmp_path / "incoming.xlsx"
    incoming.write_bytes(SAMPLE_INPUT.read_bytes())

    wb = load_workbook(incoming)
    wb[transform.SHEET_NAME]["CS9"] = "drift"
    wb.save(incoming)

    output_path = tmp_path / "actual_output.xlsx"
    exit_code = run(
        input_path=incoming,
        output_path=output_path,
        sample_path=SAMPLE_INPUT,
        report_dir=tmp_path / "reports",
        needs_review_dir=tmp_path / "needs_review",
    )

    assert exit_code == 1
    assert not output_path.exists()
    assert len(list((tmp_path / "needs_review").iterdir())) == 1
