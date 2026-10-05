from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

import transform
from transform import (
    DATE_NUMBER_FORMAT,
    TOTAL_NUMBER_FORMAT,
    data_row,
    new_subtotal_formula,
    source_for_final_column,
    translate_formula,
)


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_INPUT = ROOT / "samples" / "input.xlsx"
DATE_SOURCE = ROOT / "sh513.xlsx"


def _require_samples() -> None:
    if not SAMPLE_INPUT.exists() or not DATE_SOURCE.exists():
        pytest.skip("Put samples/input.xlsx and sh513.xlsx into the repository first.")


def _load_sheet(path: Path):
    return load_workbook(path, data_only=False)[transform.SHEET_NAME]


def _make_date_source(path: Path, value) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = transform.SHEET_NAME
    ws["B1"] = value
    wb.save(path)


@pytest.fixture(scope="module")
def transformed_output(tmp_path_factory):
    _require_samples()
    tmp_dir = tmp_path_factory.mktemp("preservation_transform")
    output_path = tmp_dir / "actual_output.xlsx"
    transform.transform(SAMPLE_INPUT, output_path, DATE_SOURCE)
    return output_path


def _expected_value(input_ws, row: int, final_col: int):
    """Value the transform should place at (row, final_col).

    Formulas are translated to their final column letters; everything else is
    copied unchanged.
    """
    source_col = source_for_final_column(final_col)
    value = input_ws.cell(row, source_col).value
    if isinstance(value, str) and value.startswith("="):
        value = translate_formula(value)
    return value


def test_normal_rows_are_preserved_across_all_regions(transformed_output):
    input_ws = _load_sheet(SAMPLE_INPUT)
    actual_ws = _load_sheet(transformed_output)
    last_data_row = data_row(input_ws)

    rows = [10, 25, 50, 75, last_data_row]
    # head (1:74), moved (75/76), shifted (77:132), post-deletion (133:162),
    # tail (163:182); 132/133 straddle the deleted column, 162/163 the tail.
    final_cols = [1, 9, 19, 74, 75, 76, 77, 100, 132, 133, 150, 162, 163, 182]

    for row in rows:
        for final_col in final_cols:
            actual = actual_ws.cell(row, final_col).value
            expected = _expected_value(input_ws, row, final_col)
            assert actual == expected, (
                f"row {row}, final col {final_col} "
                f"(source {source_for_final_column(final_col)}): "
                f"{actual!r} != {expected!r}"
            )


def test_blank_values_remain_blank_in_date_and_moved_columns(tmp_path):
    _require_samples()
    input_path = tmp_path / "input.xlsx"
    input_path.write_bytes(SAMPLE_INPUT.read_bytes())

    wb = load_workbook(input_path)
    ws = wb[transform.SHEET_NAME]
    # Blank out source positions spanning date, moved, shifted, and head columns.
    for coord in ("I10", "J10", "S10", "T10", "FF10", "BW10", "A10"):
        ws[coord] = None
    wb.save(input_path)

    date_source = tmp_path / "date_source.xlsx"
    _make_date_source(date_source, datetime(2026, 6, 30))

    output_path = tmp_path / "actual_output.xlsx"
    transform.transform(input_path, output_path, date_source)

    ws = _load_sheet(output_path)
    # Final positions: date columns stay put, FF->BW, BW->BY, A stays.
    assert ws["I10"].value is None
    assert ws["J10"].value is None
    assert ws["S10"].value is None
    assert ws["T10"].value is None
    assert ws["BW10"].value is None  # from source FF10
    assert ws["BY10"].value is None  # from source BW10
    assert ws["A10"].value is None


def test_moved_and_shifted_row7_subtotals_are_translated(transformed_output):
    input_ws = _load_sheet(SAMPLE_INPUT)
    actual_ws = _load_sheet(transformed_output)
    last_data_row = data_row(input_ws)

    # Focus on the two MOVED year-adjustment columns (75/76), the post-deletion
    # boundary (134 <- source EC), and shifted columns the existing tests do
    # not cover.
    final_cols = [75, 76, 120, 134, 150, 163]
    checked = 0
    for final_col in final_cols:
        source_col = source_for_final_column(final_col)
        source_value = input_ws.cell(7, source_col).value
        if not (isinstance(source_value, str) and "SUBTOTAL" in source_value):
            continue
        final_letter = get_column_letter(final_col)
        expected = new_subtotal_formula(final_letter, last_data_row)
        actual = actual_ws.cell(7, final_col).value
        assert actual == expected, (
            f"row-7 final col {final_col} ({final_letter}) "
            f"from source {source_col}: {actual!r} != {expected!r}"
        )
        checked += 1
    assert checked >= 2, "expected at least the two moved columns to carry SUBTOTALs"


def test_date_columns_and_new_totals_keep_display_formats(transformed_output):
    input_ws = _load_sheet(SAMPLE_INPUT)
    actual_ws = _load_sheet(transformed_output)
    last_data_row = data_row(input_ws)

    for col in (9, 10, 19, 20):  # I, J, S, T
        for row in (10, 50, last_data_row):
            assert actual_ws.cell(row, col).value == input_ws.cell(row, col).value
            assert actual_ws.cell(row, col).number_format == DATE_NUMBER_FORMAT

    assert actual_ws["BB7"].number_format == TOTAL_NUMBER_FORMAT
    assert actual_ws["CV7"].number_format == TOTAL_NUMBER_FORMAT


def test_progress_flag_does_not_change_output(tmp_path):
    _require_samples()
    out_default = tmp_path / "default.xlsx"
    out_progress = tmp_path / "progress.xlsx"
    transform.transform(SAMPLE_INPUT, out_default, DATE_SOURCE, progress=False)
    transform.transform(SAMPLE_INPUT, out_progress, DATE_SOURCE, progress=True)

    ws_default = _load_sheet(out_default)
    ws_progress = _load_sheet(out_progress)
    assert ws_default.max_column == ws_progress.max_column
    assert ws_default.max_row == ws_progress.max_row
    for row in range(1, ws_default.max_row + 1):
        for col in range(1, ws_default.max_column + 1):
            assert ws_default.cell(row, col).value == ws_progress.cell(row, col).value
