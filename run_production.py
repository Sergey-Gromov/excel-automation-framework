from __future__ import annotations

import argparse
import logging
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from excel_framework.validate_input import validate
from openpyxl import load_workbook

from transform import (
    SHEET_NAME,
    prepare_validation_pair,
    setup_logging,
    transform,
    validate_structure,
)

logger = logging.getLogger(__name__)

DEFAULT_SAMPLE = "samples/input_structural.xlsx"
DEFAULT_OUTPUT = "samples/actual_output.xlsx"
DEFAULT_REPORT_DIR = "reports"
DEFAULT_NEEDS_REVIEW_DIR = "needs_review"


def run(
    input_path: str | Path,
    output_path: str | Path = DEFAULT_OUTPUT,
    sample_path: str | Path = DEFAULT_SAMPLE,
    report_dir: str | Path = DEFAULT_REPORT_DIR,
    needs_review_dir: str | Path = DEFAULT_NEEDS_REVIEW_DIR,
    progress: bool = False,
) -> int:
    """
    Gate a production transformation run behind a structural check against
    the known-good sample. If the incoming file's structure has drifted
    (header row moved, columns added/removed, formulas/formats changed),
    quarantine it for human review instead of running the transformation —
    a silent structural mismatch is worse than a loud refusal.

    Before comparison, the input and the sample are normalized so that
    ordinary variation does not trigger quarantine: the three dated row-9
    headers are rewritten to the sample dates, and both sides are reduced to
    the smaller data volume (row-7 SUBTOTAL ranges rewritten, extra data rows
    dropped). All other structural drift is still rejected.

    Returns 0 if the transformation ran, 1 if the file was quarantined.
    """
    input_path = Path(input_path)
    sample_path = Path(sample_path)
    report_dir = Path(report_dir)
    needs_review_dir = Path(needs_review_dir)

    logger.info("Validating '%s' against sample '%s'", input_path, sample_path)
    diff_count = 0
    structural_error = None
    try:
        with tempfile.TemporaryDirectory(prefix="production_validation_") as tmp_dir:
            validation_input_path, validation_sample_path = prepare_validation_pair(
                input_path,
                sample_path,
                Path(tmp_dir),
            )
            diff_count = validate(
                validation_sample_path,
                validation_input_path,
                report_dir=report_dir,
            )
    except ValueError as exc:
        structural_error = str(exc)

    if structural_error is None:
        try:
            wb = load_workbook(input_path)
            if SHEET_NAME not in wb.sheetnames:
                structural_error = f"Лист '{SHEET_NAME}' не найден во входной книге."
            else:
                validate_structure(wb[SHEET_NAME])
        except ValueError as exc:
            structural_error = str(exc)

    if diff_count > 0 or structural_error:
        needs_review_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        case_dir = needs_review_dir / f"{stamp}_{input_path.stem}"
        case_dir.mkdir(parents=True)

        shutil.copy2(input_path, case_dir / input_path.name)
        for report_name in ("input_vs_sample.html", "input_vs_sample.json", "input_vs_sample.xlsx"):
            report_file = report_dir / report_name
            if report_file.exists():
                shutil.copy2(report_file, case_dir / report_name)

        reason = (
            f"({diff_count} difference(s))"
            if diff_count > 0
            else "structural validation failed"
        )
        print(
            f"REJECTED: '{input_path}' does not match the expected structure "
            f"{reason}. Not transformed.\n"
            + (f"Detail: {structural_error}\n" if structural_error else "")
            + f"Quarantined to: {case_dir}\n"
            f"Report: {case_dir / 'input_vs_sample.html'}",
            file=sys.stderr,
        )
        return 1

    output_path = Path(output_path)
    logger.info("Validation passed; transforming '%s' -> '%s'", input_path, output_path)
    transform(input_path, output_path, progress=progress)
    print(f"OK: '{input_path}' matches the expected structure, transformed to '{output_path}'.")
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Validate an incoming file against the sample before transforming it. "
            "Quarantines the file for human review instead of transforming it if "
            "the structure has drifted from the sample."
        )
    )
    parser.add_argument("input_path", help="Incoming file to validate and transform")
    parser.add_argument("output_path", nargs="?", default=DEFAULT_OUTPUT)
    parser.add_argument("--sample", default=DEFAULT_SAMPLE, help="Known-good reference file")
    parser.add_argument("--report-dir", default=DEFAULT_REPORT_DIR)
    parser.add_argument("--needs-review-dir", default=DEFAULT_NEEDS_REVIEW_DIR)
    parser.add_argument(
        "--progress",
        action="store_true",
        help="Force the progress bar on (default: on when stderr is a terminal)",
    )
    parser.add_argument(
        "--no-progress",
        action="store_true",
        help="Disable the progress bar",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        help="Enable debug logging",
    )
    args = parser.parse_args()

    setup_logging(args.verbose)

    if args.no_progress:
        progress = False
    elif args.progress:
        progress = True
    else:
        progress = sys.stderr.isatty()

    sys.exit(
        run(
            args.input_path,
            args.output_path,
            sample_path=args.sample,
            report_dir=args.report_dir,
            needs_review_dir=args.needs_review_dir,
            progress=progress,
        )
    )


if __name__ == "__main__":
    main()
